# -*- coding: utf-8 -*-
"""研学小结导出：优先生成 .docx，没有 python-docx 时退化为 .md。"""
from __future__ import annotations

import io
from datetime import datetime

try:
    from docx import Document
    from docx.shared import Pt, RGBColor
    HAS_DOCX = True
except Exception:                                   # pragma: no cover
    HAS_DOCX = False


def _sources_of(poi: dict, node: dict) -> list[str]:
    return [s for s in (node.get("sources") or []) if s]


def build_report(route: dict, pois: list[dict], nodes: dict, lang: str,
                 tasks: list[bool], quiz_rows: list[dict]) -> tuple[bytes, str, str]:
    """返回 (文件字节, 文件名, mime)。quiz_rows: [{q, chosen, answer, ok, explain}]"""
    en = lang == "en"
    title = ("研学小结 · " if not en else "Study Report · ") + route.get("title", "")
    stamp = datetime.now().strftime("%Y-%m-%d")
    filename = f"study-report-{route.get('id', 'route')}-{datetime.now():%Y%m%d-%H%M}.docx"

    score = sum(1 for r in quiz_rows if r["ok"])
    total = len(quiz_rows)

    if HAS_DOCX:
        doc = Document()
        style = doc.styles["Normal"]
        style.font.name = "Microsoft YaHei"
        style.font.size = Pt(10.5)
        try:                     # 中文字体需要单独指定 eastAsia
            style.element.rPr.rFonts.set(
                "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia",
                "Microsoft YaHei")
        except Exception:
            pass

        h = doc.add_heading(title, level=0)
        for run in h.runs:
            run.font.color.rgb = RGBColor(0xB1, 0x1E, 0x22)

        doc.add_paragraph(f"{'日期' if not en else 'Date'}：{stamp}")
        if route.get("summary"):
            doc.add_paragraph(route["summary"])
        if route.get("duration"):
            doc.add_paragraph(f"{'建议时长' if not en else 'Duration'}：{route['duration']}")

        doc.add_heading("一、研学路线" if not en else "1. Route", level=1)
        for i, poi in enumerate(pois, 1):
            doc.add_paragraph(f"{i}. {poi['name']}　{poi.get('addr', '')}", style="List Bullet")
            if poi.get("intro"):
                p = doc.add_paragraph(poi["intro"])
                p.paragraph_format.left_indent = Pt(16)

        doc.add_heading("二、任务卡" if not en else "2. Tasks", level=1)
        for task, done in zip(route.get("tasks", []), tasks):
            doc.add_paragraph(("☑ " if done else "☐ ") + task)

        doc.add_heading("三、随堂测评" if not en else "3. Quiz", level=1)
        doc.add_paragraph(f"{'得分' if not en else 'Score'}：{score} / {total}")
        for i, row in enumerate(quiz_rows, 1):
            mark = "✓" if row["ok"] else "✗"
            doc.add_paragraph(f"{i}. [{mark}] {row['q']}")
            doc.add_paragraph(
                f"{'你的选择' if not en else 'Your answer'}：{row['chosen']}　"
                f"{'正确答案' if not en else 'Correct'}：{row['answer']}")
            if row.get("explain"):
                p = doc.add_paragraph(row["explain"])
                p.paragraph_format.left_indent = Pt(16)

        doc.add_heading("四、资料来源" if not en else "4. Sources", level=1)
        for poi in pois:
            node = nodes.get(poi.get("entity", ""), {})
            for url in _sources_of(poi, node):
                doc.add_paragraph(f"{poi['name']}：{url}")

        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue(), filename, \
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

    # ---- 退化方案：Markdown ----
    lines = [f"# {title}", "", f"日期：{stamp}", ""]
    if route.get("summary"):
        lines += [route["summary"], ""]
    lines += ["## 一、研学路线", ""]
    for i, poi in enumerate(pois, 1):
        lines.append(f"{i}. **{poi['name']}**　{poi.get('addr', '')}")
        if poi.get("intro"):
            lines.append(f"   - {poi['intro']}")
    lines += ["", "## 二、任务卡", ""]
    for task, done in zip(route.get("tasks", []), tasks):
        lines.append(f"- [{'x' if done else ' '}] {task}")
    lines += ["", "## 三、随堂测评", "", f"得分：{score} / {total}", ""]
    for i, row in enumerate(quiz_rows, 1):
        lines.append(f"{i}. {'✓' if row['ok'] else '✗'} {row['q']}")
        lines.append(f"   - 你的选择：{row['chosen']}；正确答案：{row['answer']}")
        if row.get("explain"):
            lines.append(f"   - 解析：{row['explain']}")
    lines += ["", "## 四、资料来源", ""]
    for poi in pois:
        node = nodes.get(poi.get("entity", ""), {})
        for url in _sources_of(poi, node):
            lines.append(f"- {poi['name']}：{url}")
    data = "\n".join(lines).encode("utf-8")
    return data, filename.replace(".docx", ".md"), "text/markdown"
